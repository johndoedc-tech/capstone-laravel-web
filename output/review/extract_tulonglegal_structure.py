from pathlib import Path
from docx import Document

source = Path(r"C:\Users\JDnStefs\Downloads\Capstone 2\[FINAL PAPER] - TulongLegal.docx")
out = Path(r"C:\xampp\htdocs\capstone-laravel-web\output\review\tulonglegal-structure.txt")

doc = Document(source)
lines = []
for i, paragraph in enumerate(doc.paragraphs):
    text = " ".join(paragraph.text.split())
    if not text:
        continue
    style = paragraph.style.name if paragraph.style else ""
    lines.append(f"P{i:04d}\t[{style}]\t{text}")

lines.append("\n=== TABLES ===")
for ti, table in enumerate(doc.tables):
    lines.append(f"\nTABLE {ti}")
    for row in table.rows:
        cells = [" ".join(cell.text.split()) for cell in row.cells]
        lines.append(" | ".join(cells))

out.parent.mkdir(parents=True, exist_ok=True)
out.write_text("\n".join(lines), encoding="utf-8")
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} sections={len(doc.sections)}")
print(out)
