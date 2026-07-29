from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


def inches(value):
    return None if value is None else round(value.inches, 3)


def points(value):
    return None if value is None else round(value.pt, 2)


path = Path(sys.argv[1])
document = Document(path)

print("SECTIONS", len(document.sections))
for index, section in enumerate(document.sections, start=1):
    print(
        "SECTION",
        index,
        {
            "page_width": inches(section.page_width),
            "page_height": inches(section.page_height),
            "left_margin": inches(section.left_margin),
            "right_margin": inches(section.right_margin),
            "top_margin": inches(section.top_margin),
            "bottom_margin": inches(section.bottom_margin),
        },
    )

alignment_names = {
    None: "inherit",
    WD_ALIGN_PARAGRAPH.LEFT: "left",
    WD_ALIGN_PARAGRAPH.CENTER: "center",
    WD_ALIGN_PARAGRAPH.RIGHT: "right",
    WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
}

for index, paragraph in enumerate(document.paragraphs):
    text = " ".join(paragraph.text.split())
    if not text:
        continue
    fmt = paragraph.paragraph_format
    run_info = []
    for run in paragraph.runs:
        if not run.text.strip():
            continue
        run_info.append(
            {
                "text": " ".join(run.text.split())[:80],
                "font": run.font.name,
                "size": points(run.font.size),
                "bold": run.bold,
                "italic": run.italic,
                "underline": run.underline,
            }
        )
    print(
        f"P{index}",
        {
            "style": paragraph.style.name,
            "alignment": alignment_names.get(paragraph.alignment, str(paragraph.alignment)),
            "first_line_indent": inches(fmt.first_line_indent),
            "left_indent": inches(fmt.left_indent),
            "right_indent": inches(fmt.right_indent),
            "space_before": points(fmt.space_before),
            "space_after": points(fmt.space_after),
            "line_spacing": fmt.line_spacing,
            "keep_with_next": fmt.keep_with_next,
            "text": text,
            "runs": run_info,
        },
    )

print("TABLES", len(document.tables))
for table_index, table in enumerate(document.tables):
    print("TABLE", table_index, "rows", len(table.rows), "cols", len(table.columns))
    for row in table.rows:
        print([" ".join(cell.text.split()) for cell in row.cells])

for section_index, section in enumerate(document.sections, start=1):
    print("HEADER", section_index, [" ".join(p.text.split()) for p in section.header.paragraphs if p.text.strip()])
    print("FOOTER", section_index, [" ".join(p.text.split()) for p in section.footer.paragraphs if p.text.strip()])
