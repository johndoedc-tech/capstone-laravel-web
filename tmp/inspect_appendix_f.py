from pathlib import Path
import re

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


DOCX = Path(r"C:\xampp\htdocs\capstone-laravel-web\tmp\Harviana-appendix-f-working.docx")


def norm(text: str) -> str:
    return " ".join(text.split())


def drawings(paragraph: Paragraph) -> int:
    return len(paragraph._p.xpath(".//w:drawing")) + len(paragraph._p.xpath(".//w:pict"))


def iter_blocks(doc: Document):
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield "p", Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield "t", Table(child, doc)


doc = Document(DOCX)
paragraphs = doc.paragraphs
print("counts", len(paragraphs), len(doc.tables), len(doc.inline_shapes), len(doc.sections))

anchors = []
for i, paragraph in enumerate(paragraphs):
    text = norm(paragraph.text)
    if text in {"Features of Harviana", "Appendix F", "Supporting Interfaces of Harviana"}:
        anchors.append((i, text))
    if text.startswith("Crop Data Management") or text.startswith("User and Role Management"):
        anchors.append((i, text[:100]))
print("anchors", anchors)

for start, label in anchors:
    if label == "Features of Harviana":
        end = next(
            (i for i in range(start + 1, len(paragraphs)) if norm(paragraphs[i].text) == "Extent of the Usability of Harviana"),
            min(start + 90, len(paragraphs)),
        )
    elif label == "Appendix F":
        end = len(paragraphs)
    else:
        continue
    print("\nSECTION", label, start, end)
    for i in range(start, end):
        p = paragraphs[i]
        print(i, repr(p.text), "drawings", drawings(p), "sect", bool(p._p.xpath("./w:pPr/w:sectPr")))

print("\nFIGURE LABELS 26-39")
for i, paragraph in enumerate(paragraphs):
    text = norm(paragraph.text)
    match = re.match(r"^Figure\s+(2[6-9]|3[0-9])(?:\b|$)", text, re.IGNORECASE)
    if match:
        print(i, repr(paragraph.text), "drawings", drawings(paragraph))
