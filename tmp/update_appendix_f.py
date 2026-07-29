from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.oxml.ns import qn


DOCX = Path(r"C:\xampp\htdocs\capstone-laravel-web\tmp\Harviana-appendix-f-working.docx")
OUTPUT_DOCX = Path(r"C:\xampp\htdocs\capstone-laravel-web\tmp\Harviana-appendix-f-updated.docx")
SOURCE_DOCX = Path(r"C:\xampp\htdocs\capstone-laravel-web\tmp\Harviana-layout-working.docx")


def norm(text: str) -> str:
    return " ".join(text.split())


def drawing_count(paragraph) -> int:
    return len(paragraph._p.xpath(".//w:drawing")) + len(paragraph._p.xpath(".//w:pict"))


def copy_ppr(source, target) -> None:
    if source._p.pPr is None:
        return
    if target._p.pPr is not None:
        target._p.remove(target._p.pPr)
    target._p.insert(0, deepcopy(source._p.pPr))


def add_run_like(paragraph, text: str, source_run):
    run = paragraph.add_run(text)
    if source_run is not None and source_run._r.rPr is not None:
        run._r.insert(0, deepcopy(source_run._r.rPr))
    return run


def replace_text_like(paragraph, text: str, source_run_index: int = 0) -> None:
    source_run = None
    if paragraph.runs:
        source_run = paragraph.runs[min(source_run_index, len(paragraph.runs) - 1)]
    for child in list(paragraph._p):
        if child.tag.endswith("}r") or child.tag.endswith("}hyperlink"):
            paragraph._p.remove(child)
    add_run_like(paragraph, text, source_run)


def append_plain(doc, template, text: str, run_template_index: int = 0):
    paragraph = doc.add_paragraph()
    copy_ppr(template, paragraph)
    source_run = template.runs[min(run_template_index, len(template.runs) - 1)] if template.runs else None
    add_run_like(paragraph, text, source_run)
    return paragraph


def append_labeled(doc, template, label: str, body: str):
    paragraph = doc.add_paragraph()
    copy_ppr(template, paragraph)
    tab_run = template.runs[0] if template.runs else None
    label_run = template.runs[1] if len(template.runs) > 1 else tab_run
    body_run = template.runs[-1] if template.runs else None
    add_run_like(paragraph, "\t", tab_run)
    add_run_like(paragraph, label, label_run)
    add_run_like(paragraph, ". " + body, body_run)
    return paragraph


def extract_assets(doc, title: str):
    paragraphs = doc.paragraphs
    title_index = next(i for i, paragraph in enumerate(paragraphs) if norm(paragraph.text) == title)
    image_paragraph = next(
        paragraph for paragraph in paragraphs[title_index + 1:title_index + 5] if drawing_count(paragraph)
    )
    assets = []
    for drawing in image_paragraph._p.xpath(".//w:drawing"):
        blip = drawing.xpath(".//a:blip")[0]
        relationship_id = blip.get(qn("r:embed"))
        extent = drawing.xpath(".//wp:extent")[0]
        part = doc.part.related_parts[relationship_id]
        assets.append((part.blob, int(extent.get("cx")), int(extent.get("cy"))))
    return assets


def append_image(doc, template, assets):
    paragraph = doc.add_paragraph()
    copy_ppr(template, paragraph)
    paragraph.paragraph_format.keep_with_next = False
    paragraph.paragraph_format.keep_together = True
    for index, (blob, width, height) in enumerate(assets):
        if index:
            paragraph.add_run("  ")
        paragraph.add_run().add_picture(BytesIO(blob), width=width, height=height)
    return paragraph


def find_exact(paragraphs, text: str):
    return next(paragraph for paragraph in paragraphs if norm(paragraph.text) == text)


def find_starts(paragraphs, text: str):
    return next(paragraph for paragraph in paragraphs if norm(paragraph.text).startswith(text))


def remove_paragraph_range(paragraphs, start: int, end: int) -> None:
    for paragraph in paragraphs[start:end]:
        paragraph._element.getparent().remove(paragraph._element)


def main() -> None:
    doc = Document(DOCX)
    source = Document(SOURCE_DOCX)

    # Recover the four supporting-interface screenshots before moving content.
    crop_assets = extract_assets(doc, "Crop Data Management")
    user_assets = extract_assets(source, "User and Role Management")
    landing_assets = extract_assets(source, "Harviana Landing Page")
    login_assets = extract_assets(source, "User Login and Registration")

    paragraphs = doc.paragraphs
    feature_start = next(i for i, p in enumerate(paragraphs) if norm(p.text) == "Features of Harviana")
    feature_end = next(
        i for i in range(feature_start + 1, len(paragraphs))
        if norm(paragraphs[i].text) == "Extent of the Usability of Harviana"
    )

    # Revise the Chapter 3 overview so only core DA decision-support features remain.
    overview = paragraphs[feature_start + 1]
    for child in list(overview._p):
        if child.tag.endswith("}r") or child.tag.endswith("}hyperlink"):
            overview._p.remove(child)
    normal_run = paragraphs[feature_start + 1].runs[0] if paragraphs[feature_start + 1].runs else None
    # Use the body formatting from the Farmer Features paragraph because the
    # original overview begins with a bold tab run.
    body_template_run = paragraphs[feature_start + 2].runs[-1]
    add_run_like(
        overview,
        "\tHarviana provides role-based functions for three intended user groups: Farmers, LGU Validators, "
        "and DA Administrators. For Farmers, the system provides the Farmer Dashboard and Crop Recommendations, "
        "Crop Planning and Farmer Calendar, ML-Based Production Estimate, Damage Report Submission, Actual "
        "Harvest Recording, and access to Interactive Agricultural Map and Weather information. For LGU "
        "Validators, Harviana provides report queues, evidence review, validation notes, and the approval or "
        "return of farmer-submitted damage and harvest records. For DA Administrators, the system provides the "
        "Administrative Dashboard and Municipal Supply Forecast, Interactive Agricultural Map and Weather, and "
        "Reports and Export Generation. The following discussion presents these core features according to the "
        "responsibilities of each user group. Supporting interfaces, including Crop Data Management, User and "
        "Role Management, the landing page, and the login and registration pages, are presented in Appendix F.",
        body_template_run,
    )

    da_features = find_starts(paragraphs[feature_start:feature_end], "DA Administrator Features.")
    da_features.runs[-1].text = (
        ". The DA Administrator side of Harviana provides three core features: the Administrative Dashboard and "
        "Municipal Supply Forecast, Interactive Agricultural Map and Weather, and Reports and Export Generation. "
        "These features allow DA Administrators to monitor agricultural conditions, examine location-based "
        "production and weather information, review municipal supply estimates, and prepare reports for "
        "administrative planning and decision-making."
    )

    # Remove Crop Data Management from the core-feature discussion.
    paragraphs = doc.paragraphs
    crop_start = next(i for i, p in enumerate(paragraphs) if norm(p.text).startswith("Crop Data Management."))
    interactive_start = next(
        i for i in range(crop_start + 1, len(paragraphs))
        if norm(paragraphs[i].text).startswith("Interactive Agricultural Map and Weather.")
    )
    remove_paragraph_range(paragraphs, crop_start, interactive_start)

    # Renumber the remaining core-feature figures.
    paragraphs = doc.paragraphs
    replace_text_like(
        find_starts(paragraphs, "Figure 34 presents the Interactive Agricultural Map and Weather"),
        "\tFigure 33 presents the Interactive Agricultural Map and Weather feature, which provides Farmers and "
        "DA Administrators with a location-based view of agricultural records and current weather information.",
    )
    replace_text_like(find_exact(paragraphs, "Figure 34"), "Figure 33")
    replace_text_like(
        find_starts(paragraphs, "Figure 35 presents the Reports and Export Generation"),
        "\tFigure 34 presents the Reports and Export Generation feature, which combines agricultural summaries, "
        "planting information, actual-harvest records, and export functions for DA Administrators.",
    )
    replace_text_like(find_exact(paragraphs, "Figure 35"), "Figure 34")

    # Correct the Business Model Canvas number following the main feature figures.
    paragraphs = doc.paragraphs
    bmc_label = next(
        p for p in paragraphs
        if norm(p.text) in {"Figure 20", "Figure 37"}
        and any(norm(q.text) == "Business Model Canvas" for q in paragraphs[paragraphs.index(p) + 1:paragraphs.index(p) + 3])
    )
    replace_text_like(bmc_label, "Figure 35")
    bmc_explanation = find_starts(paragraphs, "Figure 37, showing the Business Model Canvas")
    replace_text_like(
        bmc_explanation,
        bmc_explanation.text.replace("Figure 37", "Figure 35", 1),
    )

    # Rebuild List of Figures entries 26-39; page numbers are resolved after rendering.
    paragraphs = doc.paragraphs
    lof_start = next(i for i, p in enumerate(paragraphs) if norm(p.text) == "LIST OF FIGURES")
    lot_start = next(i for i in range(lof_start + 1, len(paragraphs)) if norm(paragraphs[i].text) == "LIST OF TABLES")
    first_entry = next(i for i in range(lof_start, lot_start) if re.match(r"^26\b", norm(paragraphs[i].text)))
    entries = [
        "26\tFarmer Dashboard and Crop Recommendations . . . [PAGE]",
        "27\tCrop Planning and Farmer Calendar . . . . . . . [PAGE]",
        "28\tML-Based Production Estimate . . . . . . . . . [PAGE]",
        "29\tDamage Report Submission . . . . . . . . . . . [PAGE]",
        "30\tActual Harvest Recording . . . . . . . . . . . [PAGE]",
        "31\tLGU Report Validation . . . . . . . . . . . . . [PAGE]",
        "32\tAdministrative Dashboard and Municipal Supply Forecast . . . [PAGE]",
        "33\tInteractive Agricultural Map and Weather . . . [PAGE]",
        "34\tReports and Export Generation . . . . . . . . [PAGE]",
        "35\tBusiness Model Canvas . . . . . . . . . . . . [PAGE]",
        "36\tCrop Data Management . . . . . . . . . . . . [PAGE]",
        "37\tUser and Role Management . . . . . . . . . . [PAGE]",
        "38\tHarviana Landing Page . . . . . . . . . . . . [PAGE]",
        "39\tUser Login and Registration . . . . . . . . . [PAGE]",
    ]
    if first_entry + len(entries) > lot_start:
        raise RuntimeError("Not enough List of Figures paragraphs for entries 26-39")
    for offset, entry in enumerate(entries):
        replace_text_like(paragraphs[first_entry + offset], entry)

    # Add Appendix F to the Table of Contents using the first blank after Appendix E.
    toc_start = next(i for i, p in enumerate(paragraphs) if norm(p.text) == "TABLE OF CONTENTS")
    lof_start = next(i for i in range(toc_start + 1, len(paragraphs)) if norm(paragraphs[i].text) == "LIST OF FIGURES")
    appendix_e = next(i for i in range(toc_start, lof_start) if norm(paragraphs[i].text).startswith("E Documentation"))
    toc_f = next(i for i in range(appendix_e + 1, lof_start) if not norm(paragraphs[i].text))
    copy_ppr(paragraphs[appendix_e], paragraphs[toc_f])
    replace_text_like(paragraphs[toc_f], "F\tSupporting Interfaces of Harviana . . . . . [PAGE]")

    # Replace the Appendix F title and append its manuscript-ready content.
    paragraphs = doc.paragraphs
    appendix_f = find_starts(paragraphs, "Appendix F")
    if "Supporting Features" not in appendix_f.text:
        raise RuntimeError(f"Unexpected Appendix F title: {appendix_f.text!r}")
    appendix_f.runs[-1].text = "Supporting Interfaces of Harviana"

    labeled_template = find_starts(paragraphs, "Interactive Agricultural Map and Weather.")
    body_template = paragraphs[feature_start + 1]
    intro_template = find_starts(paragraphs, "Figure 33 presents the Interactive Agricultural Map and Weather")
    number_template = find_exact(paragraphs, "Figure 33")
    title_template = find_exact(paragraphs, "Interactive Agricultural Map and Weather")
    image_template = next(
        p for p in paragraphs[paragraphs.index(title_template) + 1:paragraphs.index(title_template) + 4]
        if drawing_count(p)
    )

    intro = append_plain(
        doc,
        body_template,
        "\tThis appendix presents the supporting administrative and access interfaces of Harviana. Crop Data "
        "Management and User and Role Management allow DA Administrators to maintain agricultural records and "
        "authorized user accounts. The Landing Page and Login and Registration interfaces introduce the platform "
        "and provide users with access to its role-based functions. Although these interfaces support the operation, "
        "security, and accessibility of Harviana, they are not included among the core decision-support features "
        "discussed in Chapter 3.",
    )

    supporting_features = [
        (
            36,
            "Crop Data Management",
            "This interface allows DA Administrators to maintain the agricultural information stored in Harviana. "
            "They can import agricultural datasets, add or update individual records, archive outdated information, "
            "and search or filter records according to municipality, crop type, production year, and other stored "
            "categories. Maintaining organized agricultural records helps ensure that the system has appropriate data "
            "for its maps, production estimates, municipal supply forecasts, dashboard summaries, and administrative reports.",
            "Figure 36 presents the Crop Data Management interface, which allows DA Administrators to maintain the "
            "agricultural records used by Harviana.",
            crop_assets,
        ),
        (
            37,
            "User and Role Management",
            "This interface allows DA Administrators to maintain the accounts authorized to access Harviana. DA "
            "Administrators can create and update user accounts, assign Farmer, LGU Validator, or DA Administrator "
            "roles, reset passwords, remove accounts, and filter users according to role or account status. These "
            "controls support the role-based access applied to the system's functions and agricultural records.",
            "Figure 37 presents the User and Role Management interface, which allows DA Administrators to maintain "
            "authorized accounts and assign the appropriate system roles.",
            user_assets,
        ),
        (
            38,
            "Harviana Landing Page",
            "The public landing page introduces the purpose of Harviana and provides visitors with general information "
            "about the platform, its features, and the development team. It serves as the initial point of access for "
            "visitors who want to learn about the system before creating an account or signing in. Links to the login "
            "and account registration pages are also available through this interface.",
            "Figure 38 presents the Harviana Landing Page, which introduces the platform as a GIS-integrated "
            "agricultural decision-support system and provides access to its public information and account pages.",
            landing_assets,
        ),
        (
            39,
            "User Login and Registration",
            "The login and registration pages provide the primary access points for Harviana's intended users. "
            "Registered users may sign in using their account credentials or Google authentication, while new Farmers "
            "may create an account by entering the required registration information. After successful authentication, "
            "Harviana directs the user to the appropriate interface according to the assigned role, such as Farmer, LGU "
            "Validator, or DA Administrator. This process helps ensure that users can access only the functions and "
            "records associated with their responsibilities.",
            "Figure 39 presents the User Login and Registration pages, which provide account access and direct "
            "authenticated users to the appropriate role-based interface.",
            login_assets,
        ),
    ]

    for number, title, body, figure_intro, assets in supporting_features:
        append_labeled(doc, labeled_template, title, body)
        p_intro = append_plain(doc, intro_template, "\t" + figure_intro)
        p_intro.paragraph_format.keep_with_next = True
        p_number = append_plain(doc, number_template, f"Figure {number}")
        p_number.paragraph_format.keep_with_next = True
        p_title = append_plain(doc, title_template, title)
        p_title.paragraph_format.keep_with_next = True
        append_image(doc, image_template, assets)

    if any("[PAGE]" in p.text for p in doc.paragraphs[lot_start:]):
        raise RuntimeError("Unexpected page placeholder outside front matter")

    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    main()
