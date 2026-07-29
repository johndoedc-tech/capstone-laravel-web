from docx import Document


SOURCE = r"C:\xampp\htdocs\capstone-laravel-web\tmp\Harviana-appendix-f-updated.docx"
OUTPUT = r"C:\xampp\htdocs\capstone-laravel-web\tmp\Harviana-appendix-f-layout.docx"


document = Document(SOURCE)

# Keep the complete Figure 34 block on one page. The two report screenshots are
# reduced only enough to fit below the introduction and captions without
# creating the large blank area seen in the source document.
for index in (603, 604, 605):
    paragraph = document.paragraphs[index]
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = True

report_images = document.paragraphs[606]
report_images.paragraph_format.keep_together = True
report_images.paragraph_format.keep_with_next = False
for inline in report_images._p.xpath(".//wp:inline"):
    extent = inline.xpath("./wp:extent")[0]
    extent.set("cx", str(round(int(extent.get("cx")) * 0.80)))
    extent.set("cy", str(round(int(extent.get("cy")) * 0.80)))
    for transform_extent in inline.xpath(".//a:xfrm/a:ext"):
        transform_extent.set("cx", extent.get("cx"))
        transform_extent.set("cy", extent.get("cy"))

# Prevent any Appendix F figure-introduction paragraph from splitting across
# pages. Captions remain chained to their corresponding screenshots.
for index in (789, 790, 791, 794, 795, 796, 799, 800, 801, 804, 805, 806):
    paragraph = document.paragraphs[index]
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = True

for index in (792, 797, 802, 807):
    paragraph = document.paragraphs[index]
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = False

document.save(OUTPUT)
print(OUTPUT)
