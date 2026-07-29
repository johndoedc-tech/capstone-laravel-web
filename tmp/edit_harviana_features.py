from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


SOURCE = Path(r"C:\Users\JDnStefs\Downloads\Harviana chap 1, 2, 3, & 4 - Copy Final.docx")
WORKING = Path(r"C:\xampp\htdocs\capstone-laravel-web\tmp\Harviana-features-working.docx")
BACKUP = Path(r"C:\xampp\htdocs\capstone-laravel-web\tmp\Harviana-features-before-edit.docx")
USER_IMAGE = Path(r"C:\xampp\htdocs\capstone-laravel-web\output\screenshots\harviana-features\06-admin-user-management.png")
CROP_IMAGE = Path(r"C:\xampp\htdocs\capstone-laravel-web\output\screenshots\harviana-features\07-admin-crop-data-management.png")


def normalized(text: str) -> str:
    return " ".join(text.split())


def paragraph_index(doc, paragraph) -> int:
    for i, candidate in enumerate(doc.paragraphs):
        if candidate._p is paragraph._p:
            return i
    raise ValueError("Paragraph is not part of the document")


def set_run_defaults(run, *, bold=None, underline=None):
    run.font.size = Pt(12)
    if bold is not None:
        run.bold = bold
    if underline is not None:
        run.underline = underline


def configure_body(paragraph, *, keep_next=False, keep_together=True):
    paragraph.style = "Body Text"
    paragraph.paragraph_format.line_spacing = 2.0
    paragraph.paragraph_format.keep_with_next = keep_next
    paragraph.paragraph_format.keep_together = keep_together
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return paragraph


def create_before(doc, reference, *, keep_next=False, keep_together=True):
    paragraph = doc.add_paragraph()
    configure_body(paragraph, keep_next=keep_next, keep_together=keep_together)
    reference.addprevious(paragraph._p)
    return paragraph


def add_plain_before(doc, reference, text: str, *, indent=True, keep_next=False):
    p = create_before(doc, reference, keep_next=keep_next)
    if indent:
        tab = p.add_run("\t")
        set_run_defaults(tab)
    r = p.add_run(text)
    set_run_defaults(r)
    return p


def add_leadin_before(doc, reference, title: str, body: str, *, keep_next=False):
    p = create_before(doc, reference, keep_next=keep_next)
    tab = p.add_run("\t")
    set_run_defaults(tab)
    lead = p.add_run(title)
    set_run_defaults(lead, underline=True)
    punct = p.add_run(". ")
    set_run_defaults(punct)
    rest = p.add_run(body)
    set_run_defaults(rest)
    return p


def add_role_summary_before(doc, reference, title: str, body: str):
    return add_leadin_before(doc, reference, title, body)


def add_figure_block_before(doc, reference, number: str, title: str, intro: str, image_elements):
    intro_p = add_plain_before(doc, reference, intro, indent=True, keep_next=True)
    intro_p.paragraph_format.keep_together = True

    num_p = add_plain_before(doc, reference, f"Figure {number}", indent=False, keep_next=True)
    title_p = add_plain_before(doc, reference, title, indent=False, keep_next=True)

    if not isinstance(image_elements, (list, tuple)):
        image_elements = [image_elements]
    for idx, el in enumerate(image_elements):
        reference.addprevious(el)
        p = next((p for p in doc.paragraphs if p._p is el), None)
        if p is not None:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.keep_together = True
            p.paragraph_format.keep_with_next = idx < len(image_elements) - 1
    return intro_p, num_p, title_p


def add_new_image_before(doc, reference, path: Path, width_inches=4.5):
    p = create_before(doc, reference, keep_next=False, keep_together=True)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_inches))
    return p._p


def append_body(doc, text: str, *, indent=True, keep_next=False):
    p = doc.add_paragraph()
    configure_body(p, keep_next=keep_next)
    if indent:
        tab = p.add_run("\t")
        set_run_defaults(tab)
    r = p.add_run(text)
    set_run_defaults(r)
    return p


def append_leadin(doc, title: str, body: str):
    p = doc.add_paragraph()
    configure_body(p)
    tab = p.add_run("\t")
    set_run_defaults(tab)
    lead = p.add_run(title)
    set_run_defaults(lead, underline=True)
    punct = p.add_run(". ")
    set_run_defaults(punct)
    rest = p.add_run(body)
    set_run_defaults(rest)
    return p


def append_caption(doc, number: str, title: str):
    num_p = append_body(doc, f"Figure {number}", indent=False, keep_next=True)
    title_p = append_body(doc, title, indent=False, keep_next=True)
    return num_p, title_p


def append_existing_element(doc, element, *, keep_next=False):
    anchor = doc.add_paragraph()
    anchor._p.addprevious(element)
    anchor._element.getparent().remove(anchor._element)
    p = next((p for p in doc.paragraphs if p._p is element), None)
    if p is not None:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.keep_together = True
        p.paragraph_format.keep_with_next = keep_next


def append_clone_paragraph(doc, template, text: str):
    p = doc.add_paragraph()
    if template._p.pPr is not None:
        p._p.insert(0, deepcopy(template._p.pPr))
    r = p.add_run(text)
    if template.runs and template.runs[0]._r.rPr is not None:
        r._r.insert(0, deepcopy(template.runs[0]._r.rPr))
    return p


def replace_paragraph_text_preserve_first_run(paragraph, text: str):
    run_rpr = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
    for child in list(paragraph._p):
        if child.tag.endswith('}r') or child.tag.endswith('}hyperlink'):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if run_rpr is not None:
        run._r.insert(0, run_rpr)


def insert_after(paragraph, new_paragraph):
    paragraph._p.addnext(new_paragraph._p)


def main():
    for path in (SOURCE, USER_IMAGE, CROP_IMAGE):
        if not path.exists():
            raise FileNotFoundError(path)

    BACKUP.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, BACKUP)
    doc = Document(SOURCE)

    feature_heading = next(p for p in doc.paragraphs if normalized(p.text).lower() == "features of harviana")
    extent_heading = next(
        p for p in doc.paragraphs
        if normalized(p.text).lower() == "extent of the usability of harviana"
    )
    feature_index = paragraph_index(doc, feature_heading)
    extent_index = paragraph_index(doc, extent_heading)
    original_feature_paragraphs = doc.paragraphs[feature_index + 1:extent_index]
    drawing_paragraphs = [
        p for p in original_feature_paragraphs
        if p._p.xpath('.//w:drawing') or p._p.xpath('.//w:pict')
    ]
    if len(drawing_paragraphs) != 12:
        raise RuntimeError(f"Expected 12 feature image paragraphs, found {len(drawing_paragraphs)}")

    (
        landing_img, login_img, farmer_img, calendar_img, ml_img, damage_img,
        harvest_img, lgu_img, admin_img, map_img, report_img_a, report_img_b,
    ) = [p._p for p in drawing_paragraphs]

    for p in original_feature_paragraphs:
        parent = p._p.getparent()
        if parent is not None:
            parent.remove(p._p)

    ref = extent_heading._p

    opening = (
        "Harviana provides role-based functions for three intended user groups: Farmers, LGU Validators, and DA Administrators. "
        "For Farmers, the system provides the Farmer Dashboard and Crop Recommendations, Crop Planning and Farmer Calendar, "
        "ML-Based Production Estimate, Damage Report Submission, Actual Harvest Recording, and access to Interactive Agricultural "
        "Map and Weather information. For LGU Validators, Harviana provides assigned report queues, evidence review, validation notes, "
        "and approval or return of farmer-submitted damage and harvest records. For DA Administrators, the system provides the "
        "Administrative Dashboard and Municipal Supply Forecast, Crop Data Management, User and Role Management, Interactive Agricultural "
        "Map and Weather, and Reports and Export Generation. The following discussion presents these core features according to the "
        "responsibilities of each user group. Supporting interfaces, including the landing page and login and registration pages, are "
        "presented in Appendix F."
    )
    add_plain_before(doc, ref, opening)

    add_role_summary_before(
        doc, ref, "Farmer Features",
        "The Farmer side of Harviana provides tools for reviewing crop information, organizing crop plans and schedules, obtaining "
        "production estimates, reporting crop damage, recording actual harvests, and accessing map and weather information."
    )

    add_leadin_before(
        doc, ref, "Farmer Dashboard and Crop Recommendations",
        "The dashboard serves as the main workspace for Farmers after login. It displays crop-related activities, active plans, reminders, "
        "harvest progress, nearby crop information, and recommendations based on the available agricultural records. Quick links are also "
        "provided for opening the calendar and other commonly used Farmer functions. This arrangement allows Farmers to review their current "
        "records before proceeding to planning, monitoring, or reporting activities."
    )
    add_figure_block_before(
        doc, ref, "26", "Farmer Dashboard and Crop Recommendations",
        "Figure 26 presents the Farmer Dashboard and Crop Recommendations feature of Harviana, which provides a summary of the Farmer's crop-related information and available actions.",
        farmer_img,
    )

    add_leadin_before(
        doc, ref, "Crop Planning and Farmer Calendar",
        "Farmers can create crop plans by providing the selected crop, intended planting area, water source, planting material, planting date, "
        "and expected harvest date. After a plan is created, its activities and important dates appear on the calendar. The page also shows "
        "active crops, upcoming reminders, and scheduled tasks such as planting, fertilizer application, crop maintenance, and harvesting. "
        "Arranging these records according to date gives Farmers a more organized way to monitor the production cycle."
    )
    add_figure_block_before(
        doc, ref, "27", "Crop Planning and Farmer Calendar",
        "Figure 27 presents the Crop Planning and Farmer Calendar feature of Harviana, where Farmers can prepare crop plans and review scheduled agricultural activities.",
        calendar_img,
    )

    add_leadin_before(
        doc, ref, "ML-Based Production Estimate",
        "During crop planning, Harviana generates an estimated production value using the information entered by the Farmer and the agricultural "
        "records processed by the machine-learning service. The estimate serves as a planning reference for the possible output of the selected "
        "crop and planted area. It is not treated as a guaranteed harvest because actual production may still be affected by weather, crop damage, "
        "farm practices, and other field conditions. The estimated value can later be compared with the actual harvest recorded for the crop plan."
    )
    add_figure_block_before(
        doc, ref, "28", "ML-Based Production Estimate",
        "Figure 28 presents the ML-Based Production Estimate displayed during the crop-planning process after the Farmer provides the required plan details.",
        ml_img,
    )

    add_leadin_before(
        doc, ref, "Damage Report Submission",
        "Farmers use this page to document crop damage that occurs during the production period. The form records the affected crop plan, cause of "
        "damage, damaged area, date of occurrence, and additional notes. Photographic evidence may also be attached to support the submitted "
        "information. Once submitted, the report receives a pending validation status and becomes available to the assigned LGU Validator. Using a "
        "standard form helps ensure that the information needed for reviewing a damage incident is recorded consistently."
    )
    add_figure_block_before(
        doc, ref, "29", "Damage Report Submission",
        "Figure 29 presents the Damage Report Submission feature of Harviana, which allows Farmers to record damage details and attach supporting evidence.",
        damage_img,
    )

    add_leadin_before(
        doc, ref, "Actual Harvest Recording",
        "After the crop-production period, Farmers can enter the actual harvest amount, unit of measurement, harvest date, optional photographic "
        "evidence, and additional notes. Harviana converts the submitted quantity into metric tons and compares it with the production estimate from "
        "the corresponding crop plan. The comparison shows the difference between the expected output and the Farmer's recorded harvest. After "
        "validation, these records contribute to production monitoring and municipal supply information."
    )
    add_figure_block_before(
        doc, ref, "30", "Actual Harvest Recording",
        "Figure 30 presents the Actual Harvest Recording feature of Harviana, where Farmers submit the production outcome of a completed crop plan.",
        harvest_img,
    )

    add_role_summary_before(
        doc, ref, "LGU Validator Features",
        "The LGU Validator side focuses on reviewing Farmer-submitted damage and actual-harvest records, examining supporting evidence, recording "
        "validation notes, and approving or returning reports for correction."
    )
    add_leadin_before(
        doc, ref, "LGU Report Validation",
        "LGU Validators can view assigned report queues, filter submissions, check validation statuses, and open the details of a damage or "
        "actual-harvest report. During the review, the validator can examine the information entered by the Farmer together with any notes and "
        "photographic evidence. A report may be approved when the information is acceptable or returned when correction or clarification is needed. "
        "Harviana records the decision, validator, validation date and time, and accompanying notes as part of the report history."
    )
    add_figure_block_before(
        doc, ref, "31", "LGU Report Validation",
        "Figure 31 presents the LGU Report Validation feature of Harviana, which provides the queue, report details, evidence, and validation actions used in reviewing Farmer submissions.",
        lgu_img,
    )

    add_role_summary_before(
        doc, ref, "DA Administrator Features",
        "The DA Administrator side provides functions for monitoring validated agricultural records, reviewing municipal supply information, managing "
        "crop data and user accounts, viewing map-based records, and preparing reports for administrative use."
    )
    add_leadin_before(
        doc, ref, "Administrative Dashboard and Municipal Supply Forecast",
        "The administrative dashboard summarizes crop records, municipalities, production information, actual harvests, reported damage, and recent "
        "system activity. It also displays municipal supply values that allow DA Administrators to review the estimated availability of selected "
        "vegetables according to location and production period. Presenting these records in one dashboard gives administrators an overview of the "
        "information collected and validated through Harviana."
    )
    add_figure_block_before(
        doc, ref, "32", "Administrative Dashboard and Municipal Supply Forecast",
        "Figure 32 presents the Administrative Dashboard and Municipal Supply Forecast of Harviana, which provide DA Administrators with a consolidated view of crop and supply-related records.",
        admin_img,
    )

    add_leadin_before(
        doc, ref, "Crop Data Management",
        "This page allows DA Administrators to import production records, add individual entries, update existing information, archive outdated records, "
        "and review summary statistics. Search and filtering options are available for locating records according to municipality, crop, year, and other "
        "stored categories. Maintaining the crop dataset is necessary because the information is used by the system's maps, production estimates, "
        "municipal supply forecasts, and administrative reports."
    )
    crop_img_el = add_new_image_before(doc, ref, CROP_IMAGE)
    # The new image was inserted first; move it after the figure introduction and caption created below.
    crop_img_el.getparent().remove(crop_img_el)
    add_figure_block_before(
        doc, ref, "33", "Crop Data Management",
        "Figure 33 presents the Crop Data Management page of Harviana, where DA Administrators maintain the agricultural records used by the system.",
        crop_img_el,
    )

    add_leadin_before(
        doc, ref, "User and Role Management",
        "DA Administrators can view registered accounts, create users, update account information, assign roles, reset passwords, and remove accounts when "
        "necessary. Accounts may be assigned as Farmer, LGU Validator, or DA Administrator according to the user's responsibilities. The page also provides "
        "search and filtering controls for locating users by role, status, or account information. These controls help maintain role-based access to the "
        "functions and records available in Harviana."
    )
    user_img_el = add_new_image_before(doc, ref, USER_IMAGE)
    user_img_el.getparent().remove(user_img_el)
    add_figure_block_before(
        doc, ref, "34", "User and Role Management",
        "Figure 34 presents the User and Role Management page of Harviana, which allows DA Administrators to maintain authorized accounts and their assigned roles.",
        user_img_el,
    )

    add_leadin_before(
        doc, ref, "Interactive Agricultural Map and Weather",
        "The map displays agricultural records according to their geographic locations. Farmers and DA Administrators can filter the information by crop "
        "type, year, view type, and farm type to focus on selected locations or production records. Weather information for a chosen area includes temperature, "
        "weather condition, rainfall probability, humidity, and wind speed. Viewing agricultural and weather information together helps users examine how "
        "crop records are distributed across the municipalities of Benguet."
    )
    add_figure_block_before(
        doc, ref, "35", "Interactive Agricultural Map and Weather",
        "Figure 35 presents the Interactive Agricultural Map and Weather feature of Harviana, which provides Farmers and DA Administrators with a location-based view of agricultural information.",
        map_img,
    )

    add_leadin_before(
        doc, ref, "Reports and Export Generation",
        "The Reports Dashboard organizes harvest accuracy, crop production, damage records, planting records, planted areas, production values, and actual "
        "harvest information into report-focused views. DA Administrators can filter the available records according to crop, municipality, production "
        "period, and report category. The Planting Report and Export page provides the corresponding detailed records and allows the information to be "
        "prepared for documentation, further analysis, and administrative reporting."
    )
    add_figure_block_before(
        doc, ref, "36", "Reports and Export Generation",
        "Figure 36 presents the Reports and Export Generation feature of Harviana, which combines summarized agricultural information with detailed planting and harvest records for export.",
        [report_img_a, report_img_b],
    )

    # Add Appendix F as a new section and reuse the original landing/login figures.
    appendix_e = next(p for p in doc.paragraphs if normalized(p.text) == "Appendix E")
    appendix_e_title = doc.paragraphs[paragraph_index(doc, appendix_e) + 1]
    new_section = doc.add_section(WD_SECTION.NEW_PAGE)
    new_section.different_first_page_header_footer = True
    new_section.header.is_linked_to_previous = True
    new_section.first_page_header.is_linked_to_previous = True
    new_section.footer.is_linked_to_previous = True
    new_section.first_page_footer.is_linked_to_previous = True

    append_clone_paragraph(doc, appendix_e, "Appendix F")
    append_clone_paragraph(doc, appendix_e_title, "Supporting Interfaces of Harviana")
    append_body(
        doc,
        "This appendix presents the supporting interfaces used to introduce Harviana and provide access to its role-based functions. These pages support "
        "navigation and account access but are not included among the core operational features discussed in Chapter 3."
    )
    append_leadin(
        doc, "Harviana Landing Page",
        "The public landing page introduces the purpose of the platform and provides visitors with general information about Harviana, its features, and "
        "the development team. Links to the login and account registration pages are also available from this interface."
    )
    append_body(
        doc,
        "Figure 38 presents the public landing page that introduces Harviana as a GIS-integrated agricultural decision-support system.",
        keep_next=True,
    )
    append_caption(doc, "38", "Harviana Landing Page")
    append_existing_element(doc, landing_img)

    append_leadin(
        doc, "User Login and Registration",
        "Registered users may sign in using their account credentials or Google authentication, while new Farmers can create an account by entering the "
        "required registration information. After authentication, Harviana directs the user to the interface associated with the assigned role."
    )
    append_body(
        doc,
        "Figure 39 presents the login and registration pages that provide the main access points for authorized users of Harviana.",
        keep_next=True,
    )
    append_caption(doc, "39", "User Login and Registration")
    append_existing_element(doc, login_img)

    # Correct the existing Appendix C Business Model Canvas caption mismatch.
    for p in doc.paragraphs:
        if normalized(p.text) == "Figure 20" and paragraph_index(doc, p) > 700:
            replace_paragraph_text_preserve_first_run(p, "Figure 37")
            break

    # Update List of Figures titles and add the two new appendix figures. Page numbers are resolved after rendering.
    lof_updates = {
        "26 Landing Page": "26 Farmer Dashboard and Crop Recommendations . . . [PAGE]",
        "27 User Login and Registration": "27 Crop Planning and Farmer Calendar . . . . . [PAGE]",
        "28 Farmer Dashboard": "28 ML-Based Production Estimate . . . . . . . . [PAGE]",
        "29 Farmer Calendar and Crop Planning": "29 Damage Report Submission . . . . . . . . . [PAGE]",
        "30 ML-Based Production Estimate": "30 Actual Harvest Recording . . . . . . . . . [PAGE]",
        "31 Damage Report Submission": "31 LGU Report Validation . . . . . . . . . . . [PAGE]",
        "32 Actual Harvest Recording": "32 Administrative Dashboard and Municipal Supply Forecast . . . [PAGE]",
        "33 LGU Report Validation": "33 Crop Data Management . . . . . . . . . . . [PAGE]",
        "34 Administrator Dashboard and Supply Forecast": "34 User and Role Management . . . . . . . . . [PAGE]",
        "35 Interactive Agricultural Map and Weather": "35 Interactive Agricultural Map and Weather . . [PAGE]",
        "36 Reports and Export Generation": "36 Reports and Export Generation . . . . . . . [PAGE]",
    }
    for p in doc.paragraphs:
        nt = normalized(p.text)
        for prefix, value in lof_updates.items():
            if nt.startswith(prefix):
                replace_paragraph_text_preserve_first_run(p, value)
                break

    fig37 = next(p for p in doc.paragraphs if normalized(p.text).startswith("37 Business Model Canvas"))
    p38 = doc.add_paragraph(style=fig37.style)
    if fig37._p.pPr is not None:
        p38._p.insert(0, deepcopy(fig37._p.pPr))
    r38 = p38.add_run("38 Harviana Landing Page . . . . . . . . . . . [PAGE]")
    if fig37.runs and fig37.runs[0]._r.rPr is not None:
        r38._r.insert(0, deepcopy(fig37.runs[0]._r.rPr))
    insert_after(fig37, p38)
    p39 = doc.add_paragraph(style=fig37.style)
    if fig37._p.pPr is not None:
        p39._p.insert(0, deepcopy(fig37._p.pPr))
    r39 = p39.add_run("39 User Login and Registration . . . . . . . . [PAGE]")
    if fig37.runs and fig37.runs[0]._r.rPr is not None:
        r39._r.insert(0, deepcopy(fig37.runs[0]._r.rPr))
    insert_after(p38, p39)

    # Add Appendix F to the manual table of contents; all affected page values are resolved after rendering.
    toc_e = next(p for p in doc.paragraphs if normalized(p.text).startswith("E Documentation"))
    toc_f = doc.add_paragraph(style=toc_e.style)
    if toc_e._p.pPr is not None:
        toc_f._p.insert(0, deepcopy(toc_e._p.pPr))
    rf = toc_f.add_run("F Supporting Interfaces of Harviana . . . . . [PAGE]")
    if toc_e.runs and toc_e.runs[0]._r.rPr is not None:
        rf._r.insert(0, deepcopy(toc_e.runs[0]._r.rPr))
    insert_after(toc_e, toc_f)

    # Remove spacer paragraphs between the final TOC entry and the section break.
    # The added Appendix F line otherwise pushes these spacers onto an almost blank page.
    candidate = toc_f._p.getnext()
    while candidate is not None:
        has_section_break = bool(candidate.xpath('./w:pPr/w:sectPr'))
        text_nodes = candidate.xpath('.//w:t')
        candidate_text = ''.join(node.text or '' for node in text_nodes).strip()
        if has_section_break or candidate_text:
            break
        following = candidate.getnext()
        candidate.getparent().remove(candidate)
        candidate = following

    WORKING.parent.mkdir(parents=True, exist_ok=True)
    doc.save(WORKING)
    print(WORKING)


if __name__ == "__main__":
    main()
