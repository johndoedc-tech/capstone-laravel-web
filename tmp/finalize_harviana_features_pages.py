from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import re

from docx import Document
from docx.oxml.ns import qn


DOCX = Path(r"C:\xampp\htdocs\capstone-laravel-web\tmp\Harviana-features-working.docx")


def norm(text: str) -> str:
    return " ".join(text.split())


def replace_text_preserve_first_run(paragraph, text: str):
    run_rpr = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
    for child in list(paragraph._p):
        if child.tag.endswith('}r') or child.tag.endswith('}hyperlink'):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if run_rpr is not None:
        run._r.insert(0, run_rpr)


def set_trailing_page(paragraph, page: int):
    text = paragraph.text
    if not re.search(r"(?:\[PAGE\]|\d+)\s*$", text):
        raise RuntimeError(f"No trailing page token found in: {text!r}")
    revised = re.sub(r"(?:\[PAGE\]|\d+)\s*$", str(page), text)
    if revised != text:
        replace_text_preserve_first_run(paragraph, revised)


def first_index(paragraphs, exact: str, start=0):
    for i in range(start, len(paragraphs)):
        if norm(paragraphs[i].text) == exact:
            return i
    raise ValueError(exact)


def main():
    doc = Document(DOCX)
    paragraphs = doc.paragraphs

    toc_start = first_index(paragraphs, "TABLE OF CONTENTS")
    lof_start = first_index(paragraphs, "LIST OF FIGURES", toc_start + 1)
    lot_start = first_index(paragraphs, "LIST OF TABLES", lof_start + 1)

    toc_pages = {
        "LIST OF TABLES": 14,
        "Background of the Study": 15,
        "Importance of the Study": 28,
        "Objectives of the Study": 31,
        "Definition of Terms": 32,
        "Software Development Methodology": 37,
        "Design Thinking": 44,
        "Scope and Delimitation": 63,
        "Data Gathering Techniques": 65,
        "Sources of Data": 67,
        "Software Development Tools": 69,
        "Information Requirements for Harviana": 75,
        "Architecture Framework of Harviana": 79,
        "Features of Harviana": 92,
        "Extent of the Usability of Harviana": 109,
        "Conclusions": 126,
        "Recommendations": 128,
        "REFERENCES": 130,
        "A Communication Letter": 137,
        "B Endorsement Letter": 138,
        "C Business Model Canvas": 139,
        "D Invitation Letter": 141,
        "E Documentation": 144,
        "F Supporting Interfaces of Harviana": 145,
    }
    for p in paragraphs[toc_start:lof_start]:
        n = norm(p.text)
        for prefix, page in toc_pages.items():
            if n.startswith(prefix) and re.search(r"(?:\[PAGE\]|\d+)\s*$", p.text):
                set_trailing_page(p, page)
                break

    figure_pages = {
        1: 39, 2: 45, 3: 46, 4: 49, 5: 51, 6: 53, 7: 53, 8: 54, 9: 56,
        10: 57, 11: 58, 12: 59, 13: 59, 14: 60, 15: 61, 16: 80, 17: 81,
        18: 83, 19: 85, 20: 86, 21: 87, 22: 88, 23: 89, 24: 90, 25: 91,
        26: 95, 27: 96, 28: 97, 29: 98, 30: 99, 31: 101, 32: 102,
        33: 104, 34: 106, 35: 107, 36: 109, 37: 139, 38: 146, 39: 147,
    }
    current = None
    for p in paragraphs[lof_start + 1:lot_start]:
        n = norm(p.text)
        match = re.match(r"^(\d+)\b", n)
        if match:
            current = int(match.group(1))
        if current in figure_pages and re.search(r"(?:\[PAGE\]|\d+)\s*$", p.text):
            set_trailing_page(p, figure_pages[current])
            current = None

    table_pages = {
        1: 78, 2: 78, 3: 78, 4: 79, 5: 111, 6: 112, 7: 114,
        8: 116, 9: 117, 10: 118, 11: 120, 12: 121, 13: 122, 14: 123,
    }
    current = None
    # The List of Tables ends at the next section break.
    for p in paragraphs[lot_start + 1:]:
        if p._p.xpath('./w:pPr/w:sectPr'):
            break
        n = norm(p.text)
        match = re.match(r"^(\d+)\b", n)
        if match:
            current = int(match.group(1))
        if current in table_pages and re.search(r"(?:\[PAGE\]|\d+)\s*$", p.text):
            set_trailing_page(p, table_pages[current])
            current = None

    # Keep Objective 3 summaries consistent with the adviser-approved core-feature list.
    for p in paragraphs:
        text = p.text
        normalized = norm(text)
        if normalized.startswith("Harviana provides role-based features for Farmers, LGU Validators, and DA Administrators.") or normalized.startswith("3. Harviana provides role-based features for Farmers, LGU Validators, and DA Administrators."):
            replace_text_preserve_first_run(
                p,
                "Harviana provides role-based features for its three user groups. "
                "Farmers can plan crops, manage calendars, estimate production, receive recommendations, report damage, "
                "record harvests, and access maps and weather. LGU Validators review and validate farmer-submitted reports. "
                "DA Administrators manage agricultural records, users and roles, municipal supply forecasts, maps, reports, "
                "and exports. These features connect Farmer submissions, LGU validation, and DA monitoring.",
            )
            run = p.runs[-1]
            run.font.name = "Courier New"
            rfonts = run._r.get_or_add_rPr().get_or_add_rFonts()
            rfonts.set(qn("w:ascii"), "Courier New")
            rfonts.set(qn("w:hAnsi"), "Courier New")
            rfonts.set(qn("w:cs"), "Courier New")
        elif "workflow for crop planning, record submission, validation, GIS-based monitoring" in text and "prediction analytics" in text:
            replace_text_preserve_first_run(p, text.replace("prediction analytics", "crop-data and user management"))
        elif normalized.startswith("3. The study identified and developed the major features of Harviana") and "prediction analytics" in text:
            replace_text_preserve_first_run(p, text.replace("prediction analytics", "user and role management"))

    if any("[PAGE]" in p.text for p in doc.paragraphs):
        raise RuntimeError("Unresolved [PAGE] placeholder remains")

    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    main()
