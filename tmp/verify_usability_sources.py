from __future__ import annotations

import csv
import statistics
import sys
from pathlib import Path

from openpyxl import load_workbook


def clean(value):
    if value is None:
        return ""
    return " ".join(str(value).split())


def inspect_xlsx(path: Path) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    print("XLSX SHEETS:", workbook.sheetnames)
    for sheet in workbook.worksheets:
        print(f"===== SHEET {sheet.title!r} rows={sheet.max_row} cols={sheet.max_column} =====")
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = [clean(value) for value in row[:40]]
            if any(values):
                print(row_index, values)
            if row_index >= 12:
                break

    sheet = workbook[workbook.sheetnames[0]]
    responses = []
    for row in sheet.iter_rows(min_row=3, values_only=True):
        if not isinstance(row[0], (int, float)):
            continue
        usefulness = [float(value) for value in row[1:9] if isinstance(value, (int, float))]
        ease_of_use = [float(value) for value in row[10:21] if isinstance(value, (int, float))]
        ease_of_learning = [float(value) for value in row[22:26] if isinstance(value, (int, float))]
        satisfaction = [float(value) for value in row[27:34] if isinstance(value, (int, float))]
        if (len(usefulness), len(ease_of_use), len(ease_of_learning), len(satisfaction)) == (8, 11, 4, 7):
            responses.append((usefulness, ease_of_use, ease_of_learning, satisfaction))
    dimension_means = []
    for dimension_index in range(4):
        values = [value for response in responses for value in response[dimension_index]]
        dimension_means.append(statistics.mean(values))
    print("DA VALID RESPONSES:", len(responses))
    print("DA DIMENSION MEANS:", [round(value, 6) for value in dimension_means])
    print("DA OVERALL (MEAN OF DIMENSIONS):", round(statistics.mean(dimension_means), 6))


def inspect_csv(path: Path) -> None:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream))
    print(f"===== CSV rows={len(rows)} cols={max(map(len, rows)) if rows else 0} =====")
    for index, row in enumerate(rows[:8], start=1):
        print(index, [clean(value) for value in row])
    valid = []
    for row in rows[1:]:
        try:
            answers = [int(row[index]) for index in range(1, 11)]
            reported_score = float(row[11])
        except (ValueError, TypeError, IndexError):
            continue
        valid.append((answers, reported_score))
    print("CSV NUMERIC RESPONSES:", len(valid))
    print("CSV REPORTED SCORE MEAN:", round(statistics.mean(score for _, score in valid), 6))
    print("CSV LAST 20:")
    for row in rows[-20:]:
        print([clean(value) for value in row])


def main() -> None:
    inspect_xlsx(Path(sys.argv[1]))
    inspect_csv(Path(sys.argv[2]))


if __name__ == "__main__":
    main()
