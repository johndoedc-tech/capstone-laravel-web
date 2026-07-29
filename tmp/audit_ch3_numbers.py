from __future__ import annotations

import csv
import json
import statistics
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


def farmer_audit(path: Path) -> dict:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream))

    responses = []
    for row in rows[1:]:
        try:
            respondent = int(row[0])
            answers = [int(row[index]) for index in range(1, 11)]
            reported_score = float(row[11])
        except (ValueError, TypeError, IndexError):
            continue
        contributions = []
        for item_number, answer in enumerate(answers, start=1):
            contributions.append(answer - 1 if item_number % 2 else 5 - answer)
        computed_score = sum(contributions) * 2.5
        responses.append(
            {
                "respondent": respondent,
                "answers": answers,
                "reported_score": reported_score,
                "computed_score": computed_score,
                "matches": abs(reported_score - computed_score) < 1e-9,
            }
        )

    item_results = []
    for item_index in range(10):
        answers = [response["answers"][item_index] for response in responses]
        counts = Counter(answers)
        raw_mean = statistics.mean(answers)
        contribution_mean = raw_mean - 1 if (item_index + 1) % 2 else 5 - raw_mean
        item_results.append(
            {
                "item": item_index + 1,
                "counts_1_to_5": [counts[value] for value in range(1, 6)],
                "raw_mean": raw_mean,
                "contribution_mean": contribution_mean,
                "contribution_mean_2dp": round(contribution_mean, 2),
            }
        )

    scores = [response["computed_score"] for response in responses]
    return {
        "valid_respondents": len(responses),
        "score_mismatches": [response["respondent"] for response in responses if not response["matches"]],
        "item_results": item_results,
        "total_score": sum(scores),
        "mean": statistics.mean(scores),
        "median": statistics.median(scores),
        "population_sd": statistics.pstdev(scores),
        "sample_sd": statistics.stdev(scores),
        "minimum": min(scores),
        "maximum": max(scores),
    }


def da_audit(path: Path) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    responses = []
    for row in sheet.iter_rows(min_row=3, max_row=20, values_only=True):
        respondent = int(row[0])
        items = []
        for column_index in list(range(1, 9)) + list(range(10, 21)) + list(range(22, 26)) + list(range(27, 34)):
            items.append(int(row[column_index]))
        responses.append({"respondent": respondent, "items": items})

    item_results = []
    for item_index in range(30):
        values = [response["items"][item_index] for response in responses]
        counts = Counter(values)
        item_results.append(
            {
                "item": item_index + 1,
                "counts_1_to_7": [counts[value] for value in range(1, 8)],
                "weighted_mean": statistics.mean(values),
                "weighted_mean_2dp": round(statistics.mean(values), 2),
            }
        )

    dimensions = {
        "Usefulness": range(0, 8),
        "Ease of Use": range(8, 19),
        "Ease of Learning": range(19, 23),
        "Satisfaction": range(23, 30),
    }
    dimension_results = {}
    for name, indices in dimensions.items():
        values = [response["items"][index] for response in responses for index in indices]
        dimension_results[name] = {
            "mean": statistics.mean(values),
            "mean_2dp": round(statistics.mean(values), 2),
        }
    exact_dimension_means = [result["mean"] for result in dimension_results.values()]
    rounded_dimension_means = [result["mean_2dp"] for result in dimension_results.values()]
    return {
        "valid_respondents": len(responses),
        "item_results": item_results,
        "dimension_results": dimension_results,
        "overall_from_exact_dimension_means": statistics.mean(exact_dimension_means),
        "overall_from_rounded_dimension_means": statistics.mean(rounded_dimension_means),
        "overall_exact_2dp": round(statistics.mean(exact_dimension_means), 2),
        "overall_rounded_2dp": round(statistics.mean(rounded_dimension_means), 2),
    }


def main() -> None:
    result = {
        "farmer": farmer_audit(Path(sys.argv[1])),
        "da": da_audit(Path(sys.argv[2])),
    }
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
