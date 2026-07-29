from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import re

from docx import Document


DOCX = Path(r"C:\xampp\htdocs\capstone-laravel-web\tmp\Harviana-layout-working.docx")


def norm(text: str) -> str:
    return " ".join(text.split())


def replace_text_preserve_first_run(paragraph, text: str) -> None:
    run_rpr = (
        deepcopy(paragraph.runs[0]._r.rPr)
        if paragraph.runs and paragraph.runs[0]._r.rPr is not None
        else None
    )
    for child in list(paragraph._p):
        if child.tag.endswith("}r") or child.tag.endswith("}hyperlink"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if run_rpr is not None:
        run._r.insert(0, run_rpr)


def set_trailing_page(paragraph, page: int) -> None:
    text = paragraph.text
    if not re.search(r"(?:\[PAGE\]|\d+)\s*$", text):
        raise RuntimeError(f"No trailing page token found in: {text!r}")
    revised = re.sub(r"(?:\[PAGE\]|\d+)\s*$", str(page), text)
    if revised != text:
        replace_text_preserve_first_run(paragraph, revised)


def first_index(paragraphs, exact: str, start: int = 0) -> int:
    for i in range(start, len(paragraphs)):
        if norm(paragraphs[i].text) == exact:
            return i
    raise ValueError(exact)


def main() -> None:
    doc = Document(DOCX)
    paragraphs = doc.paragraphs

    toc_start = first_index(paragraphs, "TABLE OF CONTENTS")
    lof_start = first_index(paragraphs, "LIST OF FIGURES", toc_start + 1)
    lot_start = first_index(paragraphs, "LIST OF TABLES", lof_start + 1)

    toc_pages = {
        "LIST OF TABLES": 14,
        "Background of the Study": 15,
        "Importance of the Study": 28,
        "Objectives of the Study": 31,
        "Definition of Terms": 32,
        "Software Development Methodology": 37,
        "Design Thinking": 44,
        "Scope and Delimitation": 63,
        "Data Gathering Techniques": 65,
        "Sources of Data": 67,
        "Software Development Tools": 69,
        "Information Requirements for Harviana": 75,
        "Architecture Framework of Harviana": 79,
        "Features of Harviana": 92,
        "Extent of the Usability of Harviana": 104,
        "Conclusions": 119,
        "Recommendations": 121,
        "REFERENCES": 123,
        "A Communication Letter": 130,
        "B Endorsement Letter": 131,
        "C Business Model Canvas": 132,
        "D Invitation Letter": 134,
        "E Documentation": 137,
        "F Supporting Interfaces of Harviana": 138,
    }
    for paragraph in paragraphs[toc_start:lof_start]:
        normalized = norm(paragraph.text)
        for prefix, page in toc_pages.items():
            if normalized.startswith(prefix) and re.search(r"(?:\[PAGE\]|\d+)\s*$", paragraph.text):
                set_trailing_page(paragraph, page)
                break

    figure_pages = {
        1: 39, 2: 45, 3: 46, 4: 49, 5: 51, 6: 53, 7: 53, 8: 54, 9: 56,
        10: 57, 11: 58, 12: 59, 13: 59, 14: 60, 15: 61, 16: 80, 17: 81,
        18: 83, 19: 85, 20: 86, 21: 87, 22: 88, 23: 89, 24: 90, 25: 91,
        26: 94, 27: 95, 28: 96, 29: 97, 30: 98, 31: 99, 32: 100,
        33: 101, 34: 102, 35: 103, 36: 104, 37: 132, 38: 138, 39: 139,
    }
    current = None
    for paragraph in paragraphs[lof_start + 1:lot_start]:
        normalized = norm(paragraph.text)
        match = re.match(r"^(\d+)\b", normalized)
        if match:
            current = int(match.group(1))
        if current in figure_pages and re.search(r"(?:\[PAGE\]|\d+)\s*$", paragraph.text):
            set_trailing_page(paragraph, figure_pages[current])
            current = None

    table_pages = {
        1: 78, 2: 78, 3: 78, 4: 79, 5: 106, 6: 107, 7: 109,
        8: 111, 9: 112, 10: 113, 11: 114, 12: 115, 13: 116, 14: 117,
    }
    current = None
    for paragraph in paragraphs[lot_start + 1:]:
        if paragraph._p.xpath("./w:pPr/w:sectPr"):
            break
        normalized = norm(paragraph.text)
        match = re.match(r"^(\d+)\b", normalized)
        if match:
            current = int(match.group(1))
        if current in table_pages and re.search(r"(?:\[PAGE\]|\d+)\s*$", paragraph.text):
            set_trailing_page(paragraph, table_pages[current])
            current = None

    if any("[PAGE]" in paragraph.text for paragraph in paragraphs):
        raise RuntimeError("Unresolved [PAGE] placeholder remains")

    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    main()
