from pathlib import Path
from docx import Document

DOCX = Path(r"C:\Users\JDnStefs\Downloads\Harviana chap 1, 2, 3, & 4 - Copy Final.docx")
doc = Document(DOCX)

feature_start = None
feature_end = None
appendix_hits = []

for i, p in enumerate(doc.paragraphs):
    text = " ".join(p.text.split())
    low = text.lower()
    if low == "features of harviana":
        feature_start = i
    if feature_start is not None and feature_end is None and low.startswith("extent of") and "usability" in low:
        feature_end = i
    if "appendix" in low:
        appendix_hits.append((i, p.style.name, text))

print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} inline_shapes={len(doc.inline_shapes)}")
print(f"feature_start={feature_start} feature_end={feature_end}")

if feature_start is not None:
    lo = max(0, feature_start - 5)
    hi = min(len(doc.paragraphs), (feature_end or feature_start + 150) + 5)
    for i in range(lo, hi):
        p = doc.paragraphs[i]
        text = " ".join(p.text.split())
        drawings = len(p._p.xpath('.//w:drawing')) + len(p._p.xpath('.//w:pict'))
        page_breaks = len(p._p.xpath('.//w:br[@w:type="page"]'))
        sect = len(p._p.xpath('./w:pPr/w:sectPr'))
        print(f"{i:04d} | {p.style.name!r} | draw={drawings} pagebr={page_breaks} sect={sect} | {text[:220]}")

print("\nAPPENDIX HITS")
for row in appendix_hits:
    print(row)

print("\nTAIL")
for i in range(max(0, len(doc.paragraphs)-120), len(doc.paragraphs)):
    p = doc.paragraphs[i]
    text = " ".join(p.text.split())
    drawings = len(p._p.xpath('.//w:drawing')) + len(p._p.xpath('.//w:pict'))
    page_breaks = len(p._p.xpath('.//w:br[@w:type="page"]'))
    sect = len(p._p.xpath('./w:pPr/w:sectPr'))
    print(f"{i:04d} | {p.style.name!r} | draw={drawings} pagebr={page_breaks} sect={sect} | {text[:220]}")
