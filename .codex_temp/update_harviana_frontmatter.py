from copy import deepcopy
from pathlib import Path
import re
import sys

from docx import Document


def replace_paragraph_text(paragraph, text: str) -> None:
    first_rpr = None
    for run in paragraph.runs:
        if run._r.rPr is not None:
            first_rpr = deepcopy(run._r.rPr)
            break
    paragraph.text = text
    if first_rpr is not None and paragraph.runs:
        paragraph.runs[0]._r.insert(0, first_rpr)


def replace_trailing_number(paragraph, number: int) -> None:
    current = paragraph.text
    if not re.search(r"\d+\s*$", current):
        raise RuntimeError(f"No trailing page number found in paragraph: {current!r}")
    updated = re.sub(r"\d+\s*$", str(number), current)
    if updated != current:
        replace_paragraph_text(paragraph, updated)


def find_unique(document: Document, starts_with: str):
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if " ".join(paragraph.text.split()).startswith(starts_with)
    ]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one paragraph beginning with {starts_with!r}, found {len(matches)}"
        )
    return matches[0]


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("Usage: update_harviana_frontmatter.py INPUT.docx OUTPUT.docx")

    source = Path(sys.argv[1])
    destination = Path(sys.argv[2])
    document = Document(source)

    # Abstract manuscript counts after the Chapter 1 expansion.
    replace_trailing_number(find_unique(document, "Total No. of Pages:"), 165)
    replace_trailing_number(find_unique(document, "1.2. Text No. of Pages:"), 120)

    # Table of Contents entries. Background remains on page 14.
    toc_pages = {
        168: 30,
        169: 34,
        170: 34,
        172: 39,
        173: 45,
        174: 64,
        175: 67,
        176: 69,
        177: 71,
        178: 75,
        180: 81,
        181: 85,
        182: 98,
        183: 110,
        185: 126,
        186: 128,
        187: 130,
        189: 138,
        190: 139,
        191: 140,
        192: 142,
        193: 145,
        194: 146,
        195: 147,
        196: 148,
        197: 149,
        198: 152,
        199: 157,
        200: 160,
    }
    for index, page in toc_pages.items():
        replace_trailing_number(document.paragraphs[index], page)

    # Make the List of Figures agree with the revised Figure 6 caption.
    replace_paragraph_text(
        document.paragraphs[208],
        "6 \tOverall Low-Fidelity Prototype of Harviana’s Main Features . . . 54",
    )

    figure_pages = {
        203: 42,
        204: 46,
        205: 48,
        206: 50,
        207: 52,
        209: 55,
        210: 56,
        211: 57,
        212: 58,
        213: 59,
        214: 60,
        215: 61,
        217: 62,
        218: 63,
        219: 86,
        220: 87,
        221: 89,
        222: 91,
        223: 92,
        225: 93,
        227: 94,
        228: 95,
        229: 96,
        230: 98,
        231: 100,
        232: 101,
        233: 102,
        234: 103,
        235: 104,
        236: 105,
        237: 106,
        238: 107,
        239: 108,
        240: 109,
        241: 140,
    }
    for index, page in figure_pages.items():
        replace_trailing_number(document.paragraphs[index], page)

    table_pages = {
        247: 84,
        248: 84,
        249: 84,
        250: 85,
        251: 111,
        252: 112,
        253: 114,
        255: 116,
        257: 117,
        259: 119,
        261: 120,
        263: 121,
        265: 123,
        266: 124,
    }
    for index, page in table_pages.items():
        replace_trailing_number(document.paragraphs[index], page)

    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)
    print(destination)


if __name__ == "__main__":
    main()
