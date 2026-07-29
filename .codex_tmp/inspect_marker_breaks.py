from __future__ import annotations

import sys

from docx import Document
from docx.oxml.ns import qn


MARKERS = [
    "LIST OF FIGURES",
    "Appendix F",
    "Appendix G",
    "Appendix H",
    "Appendix I",
    "Appendix J",
    "Appendix K",
    "Appendix L",
    "Appendix M",
    "Appendix N",
    "CURRICULUM VITAE",
    "LORD ALFREY T. BATERINA",
]


def breaks(paragraph):
    return [
        node.get(qn("w:type"), "textWrapping")
        for node in paragraph._p.xpath(".//w:br")
    ]


doc = Document(sys.argv[1])
for marker in MARKERS:
    hits = [i for i, p in enumerate(doc.paragraphs) if marker in p.text]
    index = hits[-1]
    current = doc.paragraphs[index]
    previous = doc.paragraphs[index - 1] if index else None
    print(f"MARKER={marker!r} INDEX={index}")
    print(
        " CURRENT",
        repr(current.text),
        "pageBreakBefore=",
        current.paragraph_format.page_break_before,
        "breaks=",
        breaks(current),
        "sectPr=",
        bool(current._p.xpath("./w:pPr/w:sectPr")),
    )
    if previous is not None:
        print(
            " PREVIOUS",
            repr(previous.text),
            "pageBreakBefore=",
            previous.paragraph_format.page_break_before,
            "breaks=",
            breaks(previous),
            "sectPr=",
            bool(previous._p.xpath("./w:pPr/w:sectPr")),
        )
