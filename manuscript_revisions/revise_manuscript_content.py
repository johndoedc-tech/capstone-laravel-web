from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.shared import Inches


SOURCE = Path(
    r"C:\xampp\htdocs\capstone-laravel-web\manuscript_revisions"
    r"\Harviana_Maam_Nats_Revised.docx"
)
OUTPUT = Path(
    r"C:\xampp\htdocs\capstone-laravel-web\manuscript_revisions"
    r"\Harviana_Maam_Nats_Content_Revised.docx"
)


TABLE_DESCRIPTIONS = [
    (
        0,
        "Table 1 presents the information stored for Farmer accounts in Harviana. "
        "The system automatically generates a unique user identification number for "
        "each Farmer, while the full name, email address, and encrypted password "
        "support account registration and authentication. Municipality and "
        "cooperative information associate the Farmer with the appropriate location "
        "and agricultural organization. The account status indicates whether the "
        "Farmer’s account is active or inactive and controls access to the system.",
    ),
    (
        1,
        "Table 2 presents the information stored for LGU Validator accounts. In "
        "addition to the user identification number, full name, email address, and "
        "encrypted password, the system records the municipality and barangay "
        "assigned to the LGU Validator. These location details help determine the "
        "farmer-submitted reports within the Validator’s area of responsibility. The "
        "account status controls system access, while validation notes store comments "
        "entered when reports are reviewed, approved, or returned for correction.",
    ),
    (
        2,
        "Table 3 presents the information stored for DA Administrator accounts. Each "
        "account contains an automatically generated user identification number, the "
        "Administrator’s full name, email address, and encrypted password. The "
        "user-role field identifies the account’s administrative authorization, "
        "while the account-status field determines whether access is active or "
        "inactive. These details support secure authentication and restrict "
        "administrative functions and agricultural records to authorized DA personnel.",
    ),
]


def normalized(text: str) -> str:
    return " ".join(text.split())


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    """Replace textual runs while retaining the first run's formatting."""
    first = paragraph.runs[0] if paragraph.runs else None
    formatting = None
    if first is not None:
        formatting = deepcopy(first._r.get_or_add_rPr())

    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)

    run = paragraph.add_run(text)
    if formatting is not None:
        if run._r.rPr is not None:
            run._r.remove(run._r.rPr)
        run._r.insert(0, formatting)


def clone_paragraph_after(reference: Paragraph, text: str) -> Paragraph:
    new_p = OxmlElement("w:p")
    reference._p.addnext(new_p)
    paragraph = Paragraph(new_p, reference._parent)
    if reference._p.pPr is not None:
        paragraph._p.insert(0, deepcopy(reference._p.pPr))
    run = paragraph.add_run(text)
    if reference.runs and reference.runs[0]._r.rPr is not None:
        run._r.insert(0, deepcopy(reference.runs[0]._r.rPr))
    return paragraph


def clone_paragraph_before(reference: Paragraph, text: str) -> Paragraph:
    new_p = OxmlElement("w:p")
    reference._p.addprevious(new_p)
    paragraph = Paragraph(new_p, reference._parent)
    if reference._p.pPr is not None:
        paragraph._p.insert(0, deepcopy(reference._p.pPr))
    run = paragraph.add_run(text)
    if reference.runs and reference.runs[0]._r.rPr is not None:
        run._r.insert(0, deepcopy(reference.runs[0]._r.rPr))
    return paragraph


def remove_text_only_runs(paragraph: Paragraph) -> None:
    for run in list(paragraph.runs):
        contains_visual = bool(
            run._r.xpath(
                './/*[local-name()="drawing"] | .//*[local-name()="pict"]'
            )
        )
        if not contains_visual:
            paragraph._p.remove(run._r)


def insert_table_description(table, text: str) -> None:
    new_p = OxmlElement("w:p")
    table._tbl.addnext(new_p)
    paragraph = Paragraph(new_p, table._parent)
    paragraph.style = "Body Text"
    paragraph.paragraph_format.first_line_indent = Inches(0.5)
    paragraph.paragraph_format.line_spacing = 2.0
    paragraph.paragraph_format.keep_with_next = False
    paragraph.paragraph_format.page_break_before = False
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.add_run(text)


def replace_text_everywhere(document: Document, old: str, new: str) -> int:
    count = 0
    for paragraph in document.paragraphs:
        if old in paragraph.text:
            set_paragraph_text(paragraph, paragraph.text.replace(old, new))
            count += 1
    return count


def find_paragraph(document: Document, exact_text: str) -> Paragraph:
    for paragraph in document.paragraphs:
        if normalized(paragraph.text) == exact_text:
            return paragraph
    raise ValueError(f"Paragraph not found: {exact_text!r}")


def split_label_and_title(document: Document, figure_number: int, title: str) -> None:
    combined = find_paragraph(document, f"Figure {figure_number} {title}")
    set_paragraph_text(combined, f"Figure {figure_number}")
    title_p = clone_paragraph_after(combined, title)
    title_p.paragraph_format.keep_with_next = True


def main() -> None:
    document = Document(SOURCE)

    # The panel descriptions are already present in this manuscript, but the
    # Table 2 and Table 3 labels were accidentally merged into the preceding
    # description paragraphs. Separate those labels without duplicating text.
    table1_description = next(
        p
        for p in document.paragraphs
        if normalized(p.text).startswith(
            "Table 1 presents the information stored for Farmer accounts"
        )
    )
    table1_text = table1_description.text.rstrip()
    if table1_text.endswith("Table 2"):
        set_paragraph_text(
            table1_description, table1_text[: -len("Table 2")].rstrip()
        )
        table2_label = clone_paragraph_after(table1_description, "Table 2")
        table2_label.paragraph_format.keep_with_next = True

    table2_description = next(
        p
        for p in document.paragraphs
        if normalized(p.text).startswith(
            "Table 2 presents the information stored for LGU Validator accounts"
        )
    )
    table2_text = table2_description.text.rstrip()
    if table2_text.endswith("Table 3"):
        set_paragraph_text(
            table2_description, table2_text[: -len("Table 3")].rstrip()
        )
        table3_label = clone_paragraph_after(table2_description, "Table 3")
        table3_label.paragraph_format.keep_with_next = True

    # Use the same user-role terminology throughout the manuscript.
    replace_text_everywhere(
        document, "DA-Admin Information", "DA Administrator Information"
    )
    replace_text_everywhere(
        document, "DA-Administrator Information", "DA Administrator Information"
    )

    # Correct appendix wording.
    replace_text_everywhere(document, "Relavant Code Snippets", "Relevant Code Snippets")

    # Start Chapter 4 on a fresh page.
    chapter4 = find_paragraph(document, "Chapter 4")
    chapter4.paragraph_format.page_break_before = True

    # Figure 14's label was accidentally appended to its introductory sentence.
    fig14_intro = next(
        p
        for p in document.paragraphs
        if normalized(p.text).startswith(
            "Figure 14 presents the DA Dashboard and Supply Forecast prototype"
        )
    )
    intro_text = fig14_intro.text.rstrip()
    if intro_text.endswith("Figure 14"):
        intro_text = intro_text[: -len("Figure 14")].rstrip()
    set_paragraph_text(fig14_intro, intro_text)
    fig14_label = clone_paragraph_after(fig14_intro, "Figure 14")
    fig14_label.paragraph_format.keep_with_next = True

    # Separate merged figure-number and title paragraphs.
    split_label_and_title(
        document,
        22,
        "Harviana Sequence Diagram for DA Administrator Supply Forecast",
    )
    split_label_and_title(
        document,
        23,
        "Harviana Sequence Diagram for DA Report Generation and Export",
    )
    split_label_and_title(document, 28, "ML-Based Production Estimate")
    split_label_and_title(document, 30, "Actual Harvest Recording")

    # Separate titles that share a paragraph with their image.
    for title in ["The 5 Why Analysis", "Harviana User Login and Registration Pages"]:
        paragraph = find_paragraph(document, title)
        if paragraph._p.xpath(
            './/*[local-name()="drawing"] | .//*[local-name()="pict"]'
        ):
            title_paragraph = clone_paragraph_before(paragraph, title)
            title_paragraph.paragraph_format.keep_with_next = True
            remove_text_only_runs(paragraph)

    # Correct List of Figures titles to match their actual captions.
    list_title_replacements = {
        "Harviana Farmer Report and LGU Validation":
            "Farmer Report and LGU Validation",
        "Harviana Sequence Diagram for DA Report Generation and Report":
            "Harviana Sequence Diagram for DA Report Generation and Export",
        "Farmer Dashboard": "Farmer Dashboard and Crop Recommendations",
        "Farmer Calendar and Crop Planning": "Crop Planning and Farmer Calendar",
        "Admin Dashboard and Supply Forecast":
            "Administrative Dashboard and Municipal Supply Forecast",
        "Interactive Map and Weather":
            "Interactive Agricultural Map and Weather",
    }
    for old, new in list_title_replacements.items():
        for paragraph in document.paragraphs:
            text = paragraph.text
            if old in text and text.lstrip().startswith(
                ("21", "23", "26", "27", "32", "34")
            ):
                set_paragraph_text(paragraph, text.replace(old, new))

    # Add the two appendix figures that were missing from the List of Figures.
    figure36_entry = next(
        p
        for p in document.paragraphs
        if normalized(p.text).startswith("36 Business Model Canvas")
    )
    figure37_entry = clone_paragraph_after(
        figure36_entry,
        "37\tHarviana Landing Page . . . . . . . . . . . . 000",
    )
    clone_paragraph_after(
        figure37_entry,
        "38\tHarviana User Login and Registration Pages . . . . 000",
    )

    # Keep the appendix title in the TOC consistent with the manuscript.
    replace_text_everywhere(
        document,
        "TBI Assessment Certification",
        "TBI Assessment Certificate",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
